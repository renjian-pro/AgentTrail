"""RenderStrategyAssetGenerationTest（issue #33）用的验证脚本：真正重新打开渲染产物 pptx，
按幻灯片逐张数 MSO_SHAPE_TYPE.PICTURE 类型的 shape 数量——不是假设"应该有图"，而是真的用
python-pptx 解析产物文件，证明生成的装饰性图形/封面图确实作为图片 shape 落进了每一张幻灯片。"""
import sys

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


def count_pictures(shapes):
    count = 0
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            count += 1
        elif shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            count += count_pictures(shape.shapes)
    return count


path = sys.argv[1]
prs = Presentation(path)

for index, slide in enumerate(prs.slides):
    print("slide%d_pictures=%d" % (index, count_pictures(slide.shapes)))
