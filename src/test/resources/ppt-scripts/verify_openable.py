"""PptGenerationServiceIT 用的验证脚本：用 python-pptx（这个项目选型报告里明确指定的渲染库）
重新打开一份产物 pptx，证明它不只是"看起来像"一个合法文件，而是真的能被同一个库解析、
读出幻灯片数量和标题文字。"""
import sys

from pptx import Presentation

path = sys.argv[1]
prs = Presentation(path)
slide_count = len(prs.slides)
print("slide_count=%d" % slide_count)
if slide_count >= 2:
    print("slide_count>=2")

first_slide_texts = [
    shape.text_frame.text for shape in prs.slides[0].shapes if shape.has_text_frame
]
print("first_slide_texts=%r" % first_slide_texts)
