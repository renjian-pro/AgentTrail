"""RenderStrategyTest（issue #30）用的验证脚本：真正重新打开渲染产物 pptx，读出
title_text/slide_body_text/slide_title_text 三个 shape 的文字——证明 render_ppt.py 的
apply_fill 对"模型给的内容超出 fontLimit"这个场景确实做了硬性截断，同时没超限的字段
原样保留，不是无差别覆盖。不是假设性通过，是真的用 python-pptx 读回产物文件的字符内容。"""
import sys

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


def find_shape_text(shapes, name):
    for shape in shapes:
        if shape.name == name:
            return shape.text_frame.text if shape.has_text_frame else None
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            found = find_shape_text(shape.shapes, name)
            if found is not None:
                return found
    return None


path = sys.argv[1]
prs = Presentation(path)

title_text = find_shape_text(prs.slides[0].shapes, "title_text") or ""
slide_title_text = find_shape_text(prs.slides[1].shapes, "slide_title_text") or ""
slide_body_text = find_shape_text(prs.slides[1].shapes, "slide_body_text") or ""

print("title_len=%d" % len(title_text))
print("slide_title_len=%d" % len(slide_title_text))
print("slide_body_len=%d" % len(slide_body_text))
print("title_text=%s" % title_text)
print("slide_title_text=%s" % slide_title_text)
print("slide_body_text=%s" % slide_body_text)
